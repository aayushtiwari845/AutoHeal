"""Build the editable report from implementation and measured trace evidence.

Run with the bundled Codex Python; render with packaged render_docx.py.
No live-model success rates are fabricated or inferred from scripted replay.
"""
from pathlib import Path
import json
import sys
import unittest

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from autoheal.presets import PRESETS

OUT = ROOT / "output" / "project"
QA = ROOT / "tmp" / "qa" / "report"
DOCS = ROOT / "docs"
OUT.mkdir(parents=True, exist_ok=True)
QA.mkdir(parents=True, exist_ok=True)
DOCS.mkdir(parents=True, exist_ok=True)

REFS = [
    ("ReAct", "Yao et al. ReAct Synergizing Reasoning and Acting in Language Models. ICLR 2023.", "https://arxiv.org/abs/2210.03629"),
    ("Reflexion", "Shinn et al. Reflexion Language Agents with Verbal Reinforcement Learning. 2023.", "https://arxiv.org/abs/2303.11366"),
    ("Python subprocess", "Python Software Foundation. Subprocess management documentation.", "https://docs.python.org/3/library/subprocess.html"),
    ("Ollama chat API", "Ollama. Generate a chat message. Streaming and structured response formats.", "https://docs.ollama.com/api/chat"),
    ("Qwen model card", "Qwen Team. Qwen2.5 Coder 7B Instruct official model card and license.", "https://huggingface.co/Qwen/Qwen2.5-Coder-7B-Instruct"),
    ("Qwen technical report", "Hui et al. Qwen2.5 Coder Technical Report. 2024.", "https://arxiv.org/abs/2409.12186"),
    ("Windows Job Objects", "Microsoft. Job Objects in the Windows process and thread documentation.", "https://learn.microsoft.com/en-us/windows/win32/procthread/job-objects"),
]


def text(doc, content, style=None, bold=False):
    p = doc.add_paragraph(style=style)
    p.add_run(content).bold = bold
    return p


def heading(doc, content, level=1):
    return doc.add_heading(content, level)


def bullet(doc, content):
    return text(doc, content, "List Bullet")


def code(doc, source):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1
    for index, line in enumerate(source.splitlines()):
        if index:
            p.add_run().add_break()
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for index, width in enumerate(widths):
        t.columns[index].width = Inches(width)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement("w:" + side)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "5")
        item.set(qn("w:color"), "D9D9D9")
        borders.append(item)
    t._tbl.tblPr.append(borders)
    for idx, cell in enumerate(t.rows[0].cells):
        cell.text = headers[idx]
    flag = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(flag)
    for values in rows:
        for idx, cell in enumerate(t.add_row().cells):
            cell.text = str(values[idx])
    for row_idx, row in enumerate(t.rows):
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for col_idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[col_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E7EBEE" if row_idx == 0 else "FFFFFF")
            props.append(shade)
            margin = OxmlElement("w:tcMar")
            for side in ("top", "bottom", "left", "right"):
                pad = OxmlElement("w:" + side)
                pad.set(qn("w:w"), "95")
                pad.set(qn("w:type"), "dxa")
                margin.append(pad)
            props.append(margin)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.03
                if col_idx > 0 and widths[col_idx] <= 1.1:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(10)
                    run.bold = row_idx == 0
    text(doc, "").paragraph_format.space_after = Pt(2)
    return t


def hyperlink(p, label, url):
    target = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), target)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    props.append(color)
    run.append(props)
    value = OxmlElement("w:t")
    value.text = label
    run.append(value)
    link.append(run)
    p._p.append(link)


def diagram():
    image = Image.new("RGB", (1600, 410), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 30)
    small = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 23)
    xs = [40, 355, 670, 985, 1300]
    labels = [("Generate", "Local Ollama"), ("Execute", "Child Python"), ("Evaluate", "JSON and oracle"), ("Reflect", "Bounded feedback"), ("Repair", "New candidate")]
    for index, (label, sub) in enumerate(labels):
        x = xs[index]
        draw.rectangle((x, 65, x + 260, 210), fill="#F3F5F6", outline="#9EA7AD", width=2)
        draw.text((x + 130, 116), label, fill="black", anchor="mm", font=font)
        draw.text((x + 130, 165), sub, fill="#434B50", anchor="mm", font=small)
        if index < 4:
            draw.line((x + 263, 137, xs[index + 1] - 12, 137), fill="#515B62", width=3)
            draw.polygon([(xs[index + 1] - 12, 137), (xs[index + 1] - 24, 130), (xs[index + 1] - 24, 144)], fill="#515B62")
    draw.line((800, 210, 800, 270), fill="#515B62", width=3)
    draw.text((800, 300), "PASS  stop immediately       FAIL  repair within attempt limit", anchor="mm", fill="#313A40", font=small)
    draw.line((1430, 210, 1430, 355, 170, 355, 170, 215), fill="#515B62", width=3)
    draw.polygon([(170, 215), (163, 228), (177, 228)], fill="#515B62")
    draw.text((800, 381), "Stop on success, cancellation, provider failure, or exhaustion at N <= 5", anchor="mm", fill="#313A40", font=small)
    path = QA / "state-loop.png"
    image.save(path)
    return path


def new_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.62)
    section.left_margin = section.right_margin = Inches(0.75)
    section.footer_distance = Inches(0.28)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.10
    normal.paragraph_format.space_after = Pt(7)
    doc.styles["Title"].font.size = Pt(27)
    doc.styles["Title"].font.underline = False
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    doc.styles["Title"].paragraph_format.space_after = Pt(10)
    doc.styles["Subtitle"].font.size = Pt(13)
    doc.styles["Subtitle"].paragraph_format.space_after = Pt(10)
    for name, size in (("Heading 1", 18), ("Heading 2", 12.5)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("AutoHeal Studio  |  ")
    run.font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.author = "Aayush Kumar Tiwari; Prateeksha Kini"
    doc.core_properties.title = "AutoHeal Code Execution and Self Correction"
    doc.core_properties.subject = "Generative AI lab project report"
    return doc


def build():
    summary = json.loads((ROOT / "output" / "verification" / "summary.json").read_text(encoding="utf-8"))
    if len(summary) != 9:
        raise ValueError("The original nine measured traces must be present before report authoring.")
    final_path = ROOT / "output" / "verification" / "release-summary.json"
    if not final_path.exists():
        raise ValueError("Wait for release-summary.json before final report authoring.")
    final = json.loads(final_path.read_text(encoding="utf-8"))
    if len(final) != 6:
        raise ValueError("The final six live regression traces must be present.")
    test_count = unittest.TestLoader().discover(str(ROOT / "autoheal" / "tests")).countTestCases()
    doc = new_document()

    # Page 1
    text(doc, "AutoHeal Code Execution and Self Correction", "Title")
    text(doc, "Generative AI laboratory project report", "Subtitle")
    text(doc, "Course 216U42E734  |  Semester VII  |  Academic year 2026 27")
    text(doc, "16014223003  |  Aayush Kumar Tiwari  |  LY AI & DS B2\n16014223061  |  Prateeksha Kini  |  LY AI & DS B3", bold=True)
    text(doc, "Joint submission  |  13 October 2026")
    heading(doc, "Abstract")
    text(doc, "AutoHeal Studio turns a natural-language data task into executable Python, captures failures, and asks a local coding model to repair the candidate. A browser dashboard exposes the complete trace. Three verified scenarios compare outputs with an independent reference, so successful execution alone cannot hide a logical error. Each run stops on validated success or an explicit termination condition within five attempts.")
    text(doc, "The implementation uses Python standard-library components and Ollama with Qwen2.5 Coder 7B. It includes live generation, a clearly declared controlled first failure, and a separate scripted replay. Testing demonstrates real exceptions, semantic failures, bounded retries, and cancellation. An observed live failure is retained as evidence rather than removed from the results.")
    heading(doc, "Objectives and assessment coverage")
    table(doc, ["Assessment criterion", "Marks", "Evidence in this project"], [
        ("Loop logic and state machine", "3", "Explicit phases; stdout, stderr, exit code; stop conditions and N <= 5"),
        ("Context and error reflection", "3", "Exception type and line; bounded traceback; failing JSON field paths"),
        ("Execution safety and environment", "2", "Constrained child process; five-second execution cap; Windows limits"),
        ("Viva and failure analysis", "2", "Authentic failed traces; diagnosis of syntax, runtime and logical errors"),
    ], [2.1, 0.6, 4.3])
    text(doc, "The design follows the supplied LAB CA assessment sheet. It also supplies a complete end-to-end mini-project consistent with the experiment list. The safety controls are for trusted local demonstrations and do not establish a hard operating-system sandbox.")

    # Page 2
    doc.add_page_break()
    heading(doc, "System architecture")
    text(doc, "The browser communicates with a loopback HTTP server. A background run manager calls the local Ollama service and launches a fresh Python process for each candidate. The parent process retains the trace and validates output; generated code never executes inside that parent. This separation keeps a runtime exception from crashing the dashboard.")
    p = doc.add_paragraph()
    p.add_run().add_picture(str(diagram()), width=Inches(7))
    text(doc, "Figure 1  Observable loop and explicit termination boundaries", "Caption")
    table(doc, ["Component", "Implementation", "Responsibility"], [
        ("Dashboard", "web directory", "Select task and mode; inspect candidates, feedback and results"),
        ("Local API", "server.py", "Serve UI; accept JSON runs; expose snapshots, cancellation and exports"),
        ("Orchestrator", "engine.py", "Track phases, attempts, event timestamps and terminal status"),
        ("Model transport", "provider.py", "Stream local Ollama output and bound the repair prompt"),
        ("Execution", "executor.py", "Check source; run isolated temporary child; capture and limit output"),
        ("Reference", "presets.py", "Define datasets and independent semantic validation"),
    ], [1.15, 1.2, 4.65])
    heading(doc, "States and API behavior", 2)
    text(doc, "Phases are generate, execute, evaluate, reflect and done. Terminal run states are succeeded, exhausted, cancelled and provider_error. A successful preset requires valid JSON and all expected values. A custom task requires only a non-empty JSON object and is explicitly labelled as structurally checked.")
    text(doc, "GET /api/status reports model availability and limits; GET /api/presets exposes scenarios. POST /api/runs starts an asynchronous run. GET /api/runs/{id} returns the trace snapshot; POST /api/runs/{id}/cancel stops it; GET /api/runs/{id}/export downloads JSON. The server binds only to 127.0.0.1:8765 and keeps at most 40 runs in memory, with two active runs permitted.")

    # Page 3
    doc.add_page_break()
    heading(doc, "Loop implementation and feedback")
    text(doc, "ReAct combines language-model proposals with actions and observations [1]. Reflexion uses feedback to improve a later attempt without changing model weights [2]. This project applies that general pattern to Python execution. It stores a concise technical diagnosis, not hidden model reasoning, and supplies concrete execution evidence to the next request.")
    code(doc, "for attempt in range(1, max_attempts + 1):\n    candidate = generate(task, DATA, previous_code, feedback)\n    execution = execute_in_child(candidate, timeout)\n    validation = evaluate(execution.output, independent_reference)\n    record(candidate, execution, validation)\n    if validation.passed:\n        stop_as_success()\n        break\n    feedback = bounded_error_and_validation_details(execution)\n    previous_code = candidate\nelse:\n    stop_as_exhausted()")
    text(doc, "Cancellation and provider errors stop the loop separately. A declared first failure substitutes only the first candidate. Scripted replay substitutes predetermined sources for generation while still performing actual subprocess execution and validation.")
    heading(doc, "Model contract and bounded context", 2)
    text(doc, "The local request uses Ollama POST /api/chat [4] with a JSON wrapper containing code and reflection. DATA is already a native Python value and must not be redefined or decoded again. Repair preserves chat roles: the original task and data are a user message, the previous candidate is an assistant message, and structured failure evidence is the next user message. Only the latest candidate is retained.")
    table(doc, ["Control", "Value", "Purpose"], [
        ("Model and sampling", "qwen2.5-coder:7b; temperature 0", "Local coding model; reduce sampling variation"),
        ("Context and generation", "4096 tokens; 1200 generated tokens", "Bound request and candidate size"),
        ("Previous source", "6500 characters", "Keep the latest candidate only"),
        ("Feedback", "3500 characters", "Bound structured execution evidence"),
        ("Feedback streams", "stdout 1200; stderr 1800 chars", "Retain useful evidence and suppress noise"),
        ("Provider request", "180-second bound; 10-minute keep alive", "Bound transport; avoid frequent model reloads"),
    ], [1.7, 2.3, 3.0])
    text(doc, "Feedback includes the attempt number, exception type, exact generated line when available, exit code, duration, condensed streams and failing field paths such as $.highest. A field reminder restates the task rule without reference values. Expected answers remain excluded from model repair context. Repeated lines are condensed and temporary paths sanitized.")

    # Page 4
    doc.add_page_break()
    heading(doc, "Execution boundaries and failure handling")
    text(doc, "A subprocess improves failure containment and permits enforced time limits [3]. Here, the parent parses the abstract syntax tree before launching python -I -S. The child starts in a disposable working directory, receives no stdin and inherits only an environment allowlist. DATA is injected in memory. Approved computation modules are preloaded before an audit hook denies file, network, process and native-library operations.")
    text(doc, "AST and import restrictions block disallowed modules, private attributes, dangerous builtins, introspection, dynamic evaluation and str.format traversal. Captured output is capped and compressed. JSON parsing rejects multiple outputs, duplicate keys and non-finite values. These controls provide defense in depth, not a hard OS security sandbox; hostile public execution would require a separate OS security boundary.")
    table(doc, ["Boundary", "Configured limit", "Failure behavior"], [
        ("Attempts", "1 to 5", "Explicit exhausted status after last failed candidate"),
        ("Code execution", "1 to 5 seconds", "Child killed; TimeoutError recorded"),
        ("Source and output", "24000 and 16000 characters", "Policy or output failure; capture remains bounded"),
        ("Windows Job Object", "384 MB; one active process", "Best effort; job_limit_active reported per attempt"),
        ("HTTP mutations", "Local host and origin; JSON body", "Cross-origin, rebinding-host and form posts rejected"),
    ], [1.6, 2.0, 3.4])
    heading(doc, "Failure taxonomy", 2)
    table(doc, ["Failure", "Meaning", "Recovery or termination"], [
        ("Syntax", "Source cannot parse; location available", "Usually repairable by a later candidate"),
        ("Runtime", "Candidate raises an exception", "Repair the type, operation or missing method"),
        ("Logical", "Exit code zero but answer violates rules", "Use independent validation and failing fields"),
        ("Policy or timeout", "Unsupported operation or bounded resource exceeded", "Rework computation; never remove controls to pass"),
        ("Provider", "Ollama unavailable, missing model or transport error", "Stop with provider_error; preserve evidence"),
        ("Exhaustion", "All permitted candidates failed", "Stop; human review required for a new run"),
    ], [1.1, 2.7, 3.2])
    text(doc, "Syntax errors are usually repairable. Within this run, recovery ends when attempts are exhausted, the provider fails or required capabilities are unavailable. Cancellation wins over delayed output and kills a running child. Each cancelled worker retains its capacity slot until cleanup, preventing repeated cancellation from creating unlimited background work.")

    # Page 5
    doc.add_page_break()
    heading(doc, "Scenarios and independent validation")
    heading(doc, "Sales data cleaning", 2)
    text(doc, "Eight rows mix valid numeric strings, empty and N/A amounts, and a negative adjustment. Six rows are valid. The controlled source sums raw strings and raises TypeError on line 1. A repair must parse amounts, reject invalid rows, preserve the negative value and break the regional tie alphabetically.")
    code(doc, '{"total": 520.0, "valid_rows": 6, "rejected_rows": 2,\n "by_region": {"North": 160.0, "South": 180.0, "West": 180.0},\n "top_region": "South"}')
    heading(doc, "Normalized text analytics", 2)
    text(doc, "Three documents are tokenized with contiguous ASCII letters and lowercased. The declared source calls the nonexistent lowercase() method, raising AttributeError on line 1. The correct result includes a stable top-three ranking by count and then token.")
    code(doc, '{"document_count": 3, "token_count": 14, "unique_tokens": 7,\n "frequencies": {"ai": 3, "heals": 1, "code": 4, "helps": 1,\n                 "reliable": 3, "results": 1, "tests": 1},\n "top_tokens": [{"token": "code", "count": 4},\n                {"token": "ai", "count": 3},\n                {"token": "reliable", "count": 3}]}')
    heading(doc, "Weighted grade evaluation", 2)
    text(doc, "Four students have three assessments weighted 20%, 30% and 50%. Passing requires at least 60. The seeded program uses a simple mean, exits normally, and fails semantic checks. The highest student must maximize the unrounded weighted score; alphabetic order applies only to equal greatest scores.")
    code(doc, '{"averages": {"Asha": 78.0, "Kabir": 77.0,\n              "Mira": 95.6, "Dev": 50.5},\n "class_average": 75.28, "passed": ["Asha", "Kabir", "Mira"],\n "highest": "Mira"}')
    heading(doc, "What validation establishes", 2)
    text(doc, "The validator requires the exact expected key structure, stable lists and matching finite numeric values with a small floating-point tolerance. Reference values are independently recalculated in tests using Decimal arithmetic and a separate token-counting path. The three public fixed datasets are educational cases; passing them does not prove generalization. A custom task has no reference and receives structural validation only.")

    # Page 6
    doc.add_page_break()
    heading(doc, "Measured results and verification")
    text(doc, "The original evaluation on this Windows laptop contains nine saved traces dated 13 October 2026. Three are scripted replay, three use a declared first failure followed by live model repairs, and three generate every candidate live. Durations below are complete run times, including model generation, and must not be interpreted as subprocess execution times.")
    rows = []
    for item in summary:
        label = "Replay" if item["mode"] == "replay" else "Live seeded" if item["seeded"] else "Live generated"
        rows.append((label, item["preset"].title(), item["attempts"], f"{item['duration_ms']/1000:.2f}", "Pass" if item["status"] == "succeeded" else "Exhausted"))
    table(doc, ["Mode", "Task", "Attempts", "Seconds", "Outcome"], rows, [1.65, 1.2, 1.1, 1.1, 1.95])
    text(doc, "Original outcomes were 3/3 replay, 3/3 live seeded and 2/3 fully generated live runs. The fully generated grade run exhausted five attempts. This small, deterministic set measures demonstrability and loop behavior; it is not a model benchmark or an estimate of broad repair reliability.")
    heading(doc, "Authentic failure and improvement", 2)
    text(doc, "The failing grade run first produced a malformed dataset string, then repaired the syntax but repeatedly chose min(passed_students) for highest. That returned Asha despite Mira having the greatest weighted score. A zero exit code concealed the defect until the semantic validator rejected it. Generic validation text did not localize the mistake well enough, and repeated diagnoses failed to change the offending expression.")
    text(doc, "Intermediate retests also exposed signed amounts being rejected by isdigit() and repairs that repeated an unchanged expression despite a plausible diagnosis. The revision localizes mismatched fields, includes a concise reminder of their task rule and reinforces that DATA is already loaded. Sales rules classify every row and apply alphabetical order only to a tie at the greatest regional total.")
    text(doc, "The final transport preserves assistant and user roles for the previous candidate and its failure evidence, instead of quoting both in one user prompt. This provides the model with an actual correction exchange. No reference answer values are injected. Original, intermediate and release traces remain separately available; the next page reports the newest six cases rather than selecting only successes.")

    # Page 7
    doc.add_page_break()
    heading(doc, "Final regression results")
    text(doc, "Six new live runs evaluate the final implementation: each verified task with a controlled first failure and with every candidate generated by the model. These are additional runs, not replacements for the original measurements. Durations include generation and execution on this laptop.")
    final_rows = []
    for item in final:
        final_rows.append(("Live seeded" if item["seeded"] else "Live generated", item["preset"].title(), item["attempts"], f"{item['duration_ms']/1000:.2f}", "Pass" if item["status"] == "succeeded" else item["status"].replace("_", " ").title()))
    table(doc, ["Mode", "Task", "Attempts", "Seconds", "Outcome"], final_rows, [1.65, 1.2, 1.1, 1.1, 1.95])
    passed = sum(item["status"] == "succeeded" for item in final)
    text(doc, f"The final targeted regression passed {passed}/6 live cases. Each trace preserves the actual source, error and validator decision. These six cases are too small to estimate general repair accuracy; deterministic generation, repeated prompts and public fixed datasets limit statistical interpretation.")
    heading(doc, "Feedback and convergence", 2)
    text(doc, "A useful correction identifies the failing contract, not merely an unsuccessful run. The highest-field reminder states that the name must maximize the unrounded weighted score. A rejected-row reminder states that empty and non-numeric amounts must each count once. This supports repair without giving the model a final answer to print.")
    code(doc, '{"validation": {"passed": false, "kind": "semantic",\n "mismatches": [{"path": "$.highest",\n                 "issue": "value or type differs",\n                 "rule": "Maximize weighted score; names break ties."}]}}')
    text(doc, "The schema illustration above expresses the feedback contract. The saved run traces contain the exact emitted wording. A diagnosis can still be mistaken, so the next candidate is always executed and independently evaluated before success is recorded.")
    heading(doc, "Automated verification", 2)
    text(doc, f"All {test_count} automated tests pass. They cover independent references, all repair paths, exact error lines, syntax failures, hangs, output flooding, denied access, invalid JSON, provider errors, cancellation races, first-success stopping, exhaustion, HTTP restrictions and answer-free semantic feedback. Windows Job Object activation was observed in executed local attempts.")

    # Page 8
    doc.add_page_break()
    heading(doc, "Limitations and responsible use")
    text(doc, "Qwen2.5 Coder 7B is a code-specialized instruction model [5,6]. The locally installed Ollama package uses a quantized distribution, so published checkpoint benchmark scores cannot be treated as measured performance of this application. Generation latency depends on hardware, model loading and context. Temperature zero reduces sampling variation but does not guarantee identical output across runtime versions or machines.")
    text(doc, "The Python subset excludes file analysis, internet retrieval, external packages, classes, asynchronous code and many library APIs. Timeouts can reject a correct but slow candidate. Bounded context may discard useful earlier details, and a language-model diagnosis can confidently identify the wrong cause. Independently tested hidden or varied datasets would provide stronger evidence than the current fixed examples. Custom JSON validation establishes no domain correctness.")
    heading(doc, "Privacy ethics and academic disclosure", 2)
    text(doc, "Inference calls go only to the local Ollama service. The project needs no hosted API key and does not send dataset content to a cloud model during a run. Traces stay in memory until explicitly downloaded; exports include prompts, data-derived output and source, so users should avoid sensitive personal data and review a trace before sharing it. The service must remain on loopback for trusted demonstrations.")
    text(doc, "The upstream Qwen checkpoint is licensed under Apache License 2.0 [5]. Model and distribution terms should accompany any redistribution of weights; the project archive does not bundle the downloaded model. Generated code can contain bugs or reproduce patterns with uncertain provenance, and requires review. This submission was prepared with AI-assisted implementation and writing; both students should inspect the code, understand the failure mechanics and disclose assistance according to faculty policy.")
    heading(doc, "Conclusion", 2)
    text(doc, "The completed project demonstrates code generation, execution, precise error capture, structured reflection, semantic rejection and bounded repair. Failed runs show why independent evaluation and explicit exhaustion matter. The weighted grade live case and labelled scripted replay support a clear supervisor demonstration. Other live tasks may exhaust the budget; their failures remain visible in the evidence.")
    heading(doc, "References and local evidence", 2)
    for idx, (label, description, url) in enumerate(REFS, 1):
        p = text(doc, f"[{idx}] {description} ")
        p.paragraph_format.space_after = Pt(4)
        for run in p.runs:
            run.font.size = Pt(9)
        hyperlink(p, label, url)
    text(doc, "Accessed 13 October 2026. Assessment sources: supplied LAB CA scheme and GENAI experiment list. Local evidence: output/verification/summary.json and release-summary.json, plus their corresponding JSON traces. Intermediate post-improvement and final summaries are retained separately. Implementation and tests are included in the project source.")
    doc.save(OUT / "AutoHeal_Project_Report.docx")
    print(str(OUT / "AutoHeal_Project_Report.docx"))


if __name__ == "__main__":
    build()
